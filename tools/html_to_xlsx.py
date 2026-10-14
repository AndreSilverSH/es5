#!/usr/bin/env python3
"""Create one structured Excel workbook for every HTML file in the repository.

The output directory mirrors the source directory tree. The conversion keeps
headings, paragraphs, list items, definitions, and HTML tables in document
order while dropping scripts and page navigation/chrome.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable, Iterator, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
VOID_TAGS = {
    "area", "base", "br", "col", "embed", "hr", "img", "input", "link",
    "meta", "param", "source", "track", "wbr",
}
IGNORED_TAGS = {
    "button", "head", "noscript", "script", "style", "svg", "template",
}
IGNORED_CLASSES = {"share", "share-toggler", "slogan", "erra", "rev1", "anno"}
BLOCK_TAGS = {
    "article", "blockquote", "center", "dd", "div", "dl", "dt", "fieldset",
    "figure", "footer", "form", "header", "main", "ol", "section", "ul",
}
HEADING_RE = re.compile(r"^h([1-6])$")
CONTROL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@dataclass
class Element:
    tag: str
    attrs: dict[str, str] = field(default_factory=dict)
    children: list[Element | str] = field(default_factory=list)
    parent: Element | None = None


class DocumentParser(HTMLParser):
    """Small tolerant HTML tree builder using only the Python standard library."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = Element("#document")
        self.stack = [self.root]

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        node = Element(tag, {key.lower(): value or "" for key, value in attrs})
        node.parent = self.stack[-1]
        self.stack[-1].children.append(node)
        if tag not in VOID_TAGS:
            self.stack.append(node)

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        node = Element(tag, {key.lower(): value or "" for key, value in attrs})
        node.parent = self.stack[-1]
        self.stack[-1].children.append(node)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        # HTMLParser does not perform browser-style error recovery. Closing the
        # nearest matching element is enough for the well-formed source files
        # here and also avoids corrupting the tree on older malformed fragments.
        for index in range(len(self.stack) - 1, 0, -1):
            if self.stack[index].tag == tag:
                del self.stack[index:]
                break

    def handle_data(self, data: str) -> None:
        if data:
            self.stack[-1].children.append(data)


def descendants(node: Element, tag: str | None = None) -> Iterator[Element]:
    for child in node.children:
        if isinstance(child, Element):
            if tag is None or child.tag == tag:
                yield child
            yield from descendants(child, tag)


def find_first(node: Element, tag: str) -> Element | None:
    if node.tag == tag:
        return node
    return next(descendants(node, tag), None)


def is_ignored(node: Element) -> bool:
    if node.tag in IGNORED_TAGS:
        return True
    classes = set(node.attrs.get("class", "").lower().split())
    if classes.intersection(IGNORED_CLASSES):
        return True
    # Site mastheads and dynamically populated annotation placeholders are
    # navigation chrome rather than part of the source document.
    if node.attrs.get("id", "").lower() in {
        "top", "slogan", "annotations", "timestamp", "fb-root",
    }:
        return True
    return False


def clean_text(value: str, *, preserve_lines: bool = True) -> str:
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    value = CONTROL_CHARS_RE.sub("", value)
    lines = [re.sub(r"[\t\f\v \u00a0]+", " ", line).strip() for line in value.split("\n")]
    if preserve_lines:
        # Avoid blank spacer lines introduced by nested block tags while keeping
        # meaningful <br> breaks in grammar productions and code samples.
        compact: list[str] = []
        for line in lines:
            if line or (compact and compact[-1]):
                compact.append(line)
        return "\n".join(compact).strip()
    return " ".join(part for part in lines if part).strip()


def text_content(node: Element, *, skip_tags: set[str] | None = None) -> str:
    skip_tags = skip_tags or set()
    parts: list[str] = []

    def visit(current: Element) -> None:
        if current.tag in skip_tags or is_ignored(current):
            return
        if current.tag == "br":
            parts.append("\n")
            return
        if current.tag == "img":
            alt = current.attrs.get("alt", "").strip()
            if alt:
                parts.append(alt)
            return
        if current.tag == "a":
            # These are decorative section anchors and annotation markers, not
            # prose in the source document.
            classes = set(current.attrs.get("class", "").lower().split())
            if "bak" in classes:
                return
            if clean_text("".join(text for text in raw_text_parts(current)), preserve_lines=False) == "#":
                return
        if current.tag in {"p", "div", "li", "dt", "dd", "pre"}:
            parts.append("\n")
        for child in current.children:
            if isinstance(child, Element):
                visit(child)
            else:
                parts.append(child)
        if current.tag in {"p", "div", "li", "dt", "dd", "pre"}:
            parts.append("\n")

    visit(node)
    return clean_text("".join(parts))


def raw_text_parts(node: Element) -> Iterator[str]:
    """Yield text for decorative-anchor detection without recursively rechecking it."""
    for child in node.children:
        if isinstance(child, Element):
            yield from raw_text_parts(child)
        else:
            yield child


def document_title(root: Element, source: Path) -> str:
    title_node = find_first(root, "title")
    title = text_content(title_node, skip_tags=set()) if title_node else ""
    if title:
        # The website's <title> appends annotation-key glyphs and the site name.
        title = re.split(r"\s*[|—]\s*", title, maxsplit=1)[0]
        title = re.sub(r"[#ⓉⒺ①Ⓐ]", "", title)
        title = clean_text(title, preserve_lines=False)
    if not title:
        body = find_first(root, "body") or root
        first_heading = next(
            (node for node in descendants(body) if HEADING_RE.match(node.tag)), None
        )
        title = text_content(first_heading) if first_heading else source.stem
    return title or source.stem


@dataclass
class ContentRow:
    kind: str
    section: str
    values: list[str]
    level: int = 0
    table_header: bool = False


def table_rows(table: Element) -> list[Element]:
    found: list[Element] = []

    def visit(node: Element) -> None:
        for child in node.children:
            if not isinstance(child, Element):
                continue
            if child.tag == "table":
                continue
            if child.tag == "tr":
                found.append(child)
            else:
                visit(child)

    visit(table)
    return found


def table_cells(row: Element) -> list[Element]:
    cells: list[Element] = []

    def visit(node: Element) -> None:
        for child in node.children:
            if not isinstance(child, Element):
                continue
            if child.tag in {"td", "th"}:
                cells.append(child)
            else:
                visit(child)

    visit(row)
    return cells


def table_caption(table: Element) -> str:
    caption = next(descendants(table, "caption"), None)
    return text_content(caption, skip_tags={"table"}) if caption else ""


def is_header_row(row: Element, cells: Sequence[Element], row_index: int) -> bool:
    if any(cell.tag == "th" for cell in cells):
        return True
    if row_index == 0 and any(
        "table-header" in cell.attrs.get("class", "").lower().split()
        or cell.attrs.get("bgcolor", "").lower() in {"#c0c0c0", "silver"}
        for cell in cells
    ):
        return True
    return False


def extract_rows(root: Element) -> list[ContentRow]:
    body = find_first(root, "body") or root
    output: list[ContentRow] = []
    section_stack: list[tuple[int, str]] = []
    table_number = 0

    def current_section() -> str:
        return " › ".join(title for _, title in section_stack)

    def emit(kind: str, values: Sequence[str], *, level: int = 0, header: bool = False) -> None:
        cleaned = [clean_text(value) for value in values]
        if any(cleaned):
            output.append(ContentRow(kind, current_section(), cleaned, level, header))

    def process_table(table: Element) -> None:
        nonlocal table_number
        table_number += 1
        caption = table_caption(table)
        if caption:
            emit(f"Подпись таблицы {table_number}", [caption])
        rows = table_rows(table)
        for index, row in enumerate(rows, start=1):
            cells = table_cells(row)
            values = [text_content(cell, skip_tags={"table"}) for cell in cells]
            if not any(values):
                continue
            header = is_header_row(row, cells, index - 1)
            kind = f"Таблица {table_number} · заголовок" if header else f"Таблица {table_number} · строка {index}"
            emit(kind, values, header=header)

    def process_list(list_node: Element, depth: int) -> None:
        list_kind = "Нумерованный список" if list_node.tag == "ol" else "Маркированный список"
        for child in list_node.children:
            if not isinstance(child, Element):
                continue
            if child.tag == "li":
                emit(list_kind, [text_content(child, skip_tags={"ul", "ol"})], level=depth)
                for nested in child.children:
                    if isinstance(nested, Element) and nested.tag in {"ul", "ol"}:
                        process_list(nested, depth + 1)
            elif child.tag in {"ul", "ol"}:
                process_list(child, depth + 1)
            else:
                visit_container(child, depth)

    def visit_container(container: Element, list_depth: int = 0, inside_nav: bool = False) -> None:
        pending_inline: list[str] = []

        def flush_inline() -> None:
            text = clean_text("".join(pending_inline))
            if text:
                emit("Текст", [text])
            pending_inline.clear()

        for child in container.children:
            if isinstance(child, str):
                if not inside_nav:
                    pending_inline.append(child)
                continue
            if is_ignored(child):
                continue
            classes = set(child.attrs.get("class", "").lower().split())
            if inside_nav and (child.tag == "a" or "toc" in classes):
                continue
            if child.tag == "nav":
                flush_inline()
                visit_container(child, list_depth, True)
                continue
            heading_match = HEADING_RE.match(child.tag)
            if heading_match:
                flush_inline()
                level = int(heading_match.group(1))
                heading = text_content(child)
                if heading:
                    while section_stack and section_stack[-1][0] >= level:
                        section_stack.pop()
                    section_stack.append((level, heading))
                    emit(f"Заголовок H{level}", [heading], level=level)
                continue
            if child.tag == "table":
                flush_inline()
                process_table(child)
                continue
            if child.tag in {"ul", "ol"}:
                flush_inline()
                process_list(child, list_depth)
                continue
            if child.tag == "li":
                flush_inline()
                emit("Пункт списка", [text_content(child, skip_tags={"ul", "ol"})], level=list_depth)
                for nested in child.children:
                    if isinstance(nested, Element) and nested.tag in {"ul", "ol"}:
                        process_list(nested, list_depth + 1)
                continue
            if child.tag in {"p", "pre", "blockquote", "dt", "dd"}:
                flush_inline()
                if child.tag == "pre":
                    kind = "Код / пример"
                elif child.tag == "blockquote":
                    kind = "Цитата"
                elif child.tag == "dt":
                    kind = "Термин"
                elif child.tag == "dd":
                    kind = "Определение"
                else:
                    kind = "Абзац"
                emit(kind, [text_content(child)])
                continue
            if child.tag in {"br", "img"}:
                if child.tag == "br":
                    pending_inline.append("\n")
                else:
                    alt = child.attrs.get("alt", "").strip()
                    if alt:
                        pending_inline.append(alt)
                continue
            if child.tag == "hr":
                flush_inline()
                continue
            # Inline tags belong to the surrounding text. Block containers are
            # traversed so their semantic children retain document order.
            if child.tag in BLOCK_TAGS or any(
                isinstance(grandchild, Element)
                and (HEADING_RE.match(grandchild.tag) or grandchild.tag in {
                    "p", "pre", "blockquote", "dt", "dd", "table", "ul", "ol", "li",
                })
                for grandchild in child.children
            ):
                flush_inline()
                visit_container(child, list_depth, inside_nav or child.tag == "nav")
            else:
                pending_inline.append(text_content(child))
        flush_inline()

    visit_container(body)
    return output


def safe_cell(value: str) -> str:
    value = CONTROL_CHARS_RE.sub("", value)
    # Excel limits a cell to 32,767 characters. Preserve the beginning and end
    # with an explicit marker if an unusually large source paragraph exceeds it.
    if len(value) > 32767:
        marker = "\n[… текст сокращён: ограничение Excel 32 767 символов …]\n"
        available = 32767 - len(marker)
        value = value[: available // 2] + marker + value[-(available - available // 2):]
    return value


def set_literal(cell, value: str) -> None:
    cell.value = safe_cell(value)
    # Source documents can contain strings beginning with '='; keep every value
    # literal so the export cannot accidentally create or evaluate formulas.
    if isinstance(cell.value, str):
        cell.data_type = "s"


def make_workbook(relative_path: Path, rows: Sequence[ContentRow], title: str):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Содержимое"

    max_content_columns = max((len(row.values) for row in rows), default=1)
    total_columns = max(4, 3 + max_content_columns)
    last_column = get_column_letter(total_columns)

    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)
    set_literal(worksheet.cell(1, 1), title)
    worksheet.cell(1, 1).font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    worksheet.cell(1, 1).fill = PatternFill("solid", fgColor="17365D")
    worksheet.cell(1, 1).alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 32

    set_literal(worksheet.cell(2, 1), "Источник HTML")
    set_literal(worksheet.cell(2, 2), relative_path.as_posix())
    worksheet.cell(2, 1).font = Font(name="Arial", bold=True, color="17365D")
    worksheet.merge_cells(start_row=2, start_column=2, end_row=2, end_column=total_columns)
    worksheet.cell(2, 2).alignment = Alignment(wrap_text=True)

    set_literal(worksheet.cell(3, 1), "Элементов")
    set_literal(worksheet.cell(3, 2), str(len(rows)))
    worksheet.cell(3, 1).font = Font(name="Arial", bold=True, color="17365D")

    header_row = 5
    headings = ["№", "Тип", "Раздел", "Содержимое"]
    headings.extend(f"Ячейка {index}" for index in range(2, total_columns - 2))
    for column, heading in enumerate(headings, start=1):
        cell = worksheet.cell(header_row, column)
        set_literal(cell, heading)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24527A")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[header_row].height = 25

    thin_gray = Side(style="hair", color="D9E2F3")
    for row_index, item in enumerate(rows, start=header_row + 1):
        values = [str(row_index - header_row), item.kind, item.section, *item.values]
        for column in range(1, total_columns + 1):
            value = values[column - 1] if column <= len(values) else ""
            cell = worksheet.cell(row_index, column)
            set_literal(cell, value)
            cell.font = Font(name="Arial", size=10, color="1F1F1F")
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
                indent=min(item.level, 15) if column == 4 else 0,
            )
            cell.border = Border(bottom=thin_gray)

        kind_cell = worksheet.cell(row_index, 2)
        if item.kind.startswith("Заголовок"):
            level = min(max(item.level, 1), 6)
            fills = ["D9EAF7", "EAF2F8", "F2F6FA", "F7F9FC", "FBFCFD", "FFFFFF"]
            kind_cell.fill = PatternFill("solid", fgColor=fills[level - 1])
            kind_cell.font = Font(name="Arial", size=max(10, 14 - level), bold=True, color="17365D")
            worksheet.cell(row_index, 3).fill = PatternFill("solid", fgColor=fills[level - 1])
            worksheet.cell(row_index, 4).fill = PatternFill("solid", fgColor=fills[level - 1])
        elif item.table_header:
            for column in range(2, total_columns + 1):
                cell = worksheet.cell(row_index, column)
                cell.fill = PatternFill("solid", fgColor="D9EAD3")
                cell.font = Font(name="Arial", size=10, bold=True, color="274E13")
        elif item.kind.startswith("Подпись таблицы"):
            kind_cell.font = Font(name="Arial", size=10, bold=True, italic=True, color="17365D")
        elif item.kind in {"Маркированный список", "Нумерованный список", "Пункт списка"}:
            kind_cell.font = Font(name="Arial", size=10, color="38761D")

    last_row = max(header_row, worksheet.max_row)
    worksheet.freeze_panes = "A6"
    worksheet.auto_filter.ref = f"A{header_row}:{last_column}{last_row}"
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 7
    worksheet.column_dimensions["B"].width = 26
    worksheet.column_dimensions["C"].width = 42
    worksheet.column_dimensions["D"].width = 90
    for column in range(5, total_columns + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 44

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_title_rows = "1:5"
    worksheet.sheet_properties.pageSetUpPr.autoPageBreaks = False
    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.page_margins.top = 0.5
    worksheet.page_margins.bottom = 0.5

    workbook.properties.title = title[:255]
    workbook.properties.subject = f"Экспорт из {relative_path.as_posix()}"
    workbook.properties.creator = "HTML → XLSX export"
    return workbook


def parse_html(path: Path) -> tuple[str, list[ContentRow]]:
    source_text = path.read_text(encoding="utf-8", errors="replace")
    parser = DocumentParser()
    parser.feed(source_text)
    parser.close()
    title = document_title(parser.root, path)
    return title, extract_rows(parser.root)


def is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def find_html_files(source_dir: Path, output_dir: Path) -> list[Path]:
    excluded_names = {".git", ".venv", "node_modules", "__pycache__", "build", "dist", "coverage"}
    files: list[Path] = []
    output_resolved = output_dir.resolve()
    for path in source_dir.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in {".html", ".htm"}:
            continue
        resolved = path.resolve()
        if is_relative_to(resolved, output_resolved):
            continue
        relative = path.relative_to(source_dir)
        if excluded_names.intersection(relative.parts):
            continue
        files.append(path)
    return sorted(files, key=lambda path: path.relative_to(source_dir).as_posix().lower())


def section_heading_index(source_dir: Path, sources: Sequence[Path]) -> dict[str, str]:
    """Map section IDs to their visible headings in the top-level spec pages."""
    headings: dict[str, str] = {}
    for source in sources:
        if source.parent != source_dir:
            continue
        parser = DocumentParser()
        parser.feed(source.read_text(encoding="utf-8", errors="replace"))
        parser.close()
        for node in descendants(parser.root):
            identifier = node.attrs.get("id", "").strip()
            if identifier and HEADING_RE.match(node.tag):
                heading = text_content(node)
                if heading:
                    headings.setdefault(identifier, heading)
    return headings


def workbook_title(relative: Path, html_title: str, section_headings: dict[str, str]) -> str:
    """Use page titles for documents and section context for HTML fragments."""
    prefixes = {
        "anno": "Аннотация",
        "erra": "Исправление",
        "rev1": "Редакция 5.1",
    }
    category = relative.parts[0] if len(relative.parts) > 1 else ""
    if category not in prefixes:
        return html_title

    section_id = relative.stem
    section_title = section_headings.get(section_id, "")
    if section_title:
        visible_id = section_id[1:] if section_id.startswith("x") else section_id
        if visible_id and visible_id not in section_title:
            section_title = f"{visible_id} {section_title}"
        pieces = [prefixes[category], section_title]
    else:
        pieces = [prefixes[category]]

    if html_title and html_title not in {section_id, section_title}:
        pieces.append(html_title)
    if len(pieces) == 1:
        pieces.append(section_id)
    return " — ".join(pieces)


def safe_filename_stem(value: str, fallback: str) -> str:
    value = unicodedata.normalize("NFKC", value)
    value = re.sub(r"[^\w.-]+", "_", value, flags=re.UNICODE)
    value = re.sub(r"_+", "_", value).strip(" ._-")
    if not value:
        value = re.sub(r"[^\w.-]+", "_", fallback, flags=re.UNICODE).strip(" ._-") or "document"
    # Keep filenames portable to Windows and reasonably short for archives.
    if len(value) > 145:
        shortened = value[:145]
        boundary = shortened.rfind("_")
        value = shortened[:boundary] if boundary >= 110 else shortened
    value = value.rstrip(" ._-") or "document"
    if re.match(r"^(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\.|$)", value, flags=re.IGNORECASE):
        value = f"_{value}"
    return value


def output_relative_path(relative: Path, title: str, used: set[str]) -> Path:
    stem = safe_filename_stem(title, relative.stem)
    candidate = relative.parent / f"{stem}.xlsx"
    key = candidate.as_posix().casefold()
    if key in used:
        source_suffix = safe_filename_stem(relative.stem, "source")
        stem = safe_filename_stem(f"{stem}_{source_suffix}", "document")
        candidate = relative.parent / f"{stem}.xlsx"
        key = candidate.as_posix().casefold()
        counter = 2
        while key in used:
            candidate = relative.parent / f"{stem}_{counter}.xlsx"
            key = candidate.as_posix().casefold()
            counter += 1
    used.add(key)
    return candidate


def previous_manifest_paths(manifest_path: Path) -> set[Path]:
    if not manifest_path.exists():
        return set()
    with manifest_path.open("r", encoding="utf-8-sig", newline="") as manifest_file:
        return {
            Path(row["xlsx_file"])
            for row in csv.DictReader(manifest_file, delimiter=";")
            if row.get("xlsx_file")
        }


def write_manifest(manifest_path: Path, rows: Sequence[dict[str, str]]) -> None:
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    with manifest_path.open("w", encoding="utf-8-sig", newline="") as manifest_file:
        writer = csv.DictWriter(
            manifest_file,
            fieldnames=["source_html", "xlsx_file", "title"],
            delimiter=";",
        )
        writer.writeheader()
        writer.writerows(rows)


def convert(source_dir: Path, output_dir: Path, files: Iterable[Path] | None = None) -> int:
    source_dir = source_dir.resolve()
    output_dir = output_dir.resolve()
    sources = list(files) if files is not None else find_html_files(source_dir, output_dir)
    if not sources:
        raise RuntimeError(f"В каталоге {source_dir} не найдено HTML-файлов")

    section_headings = section_heading_index(source_dir, sources)
    manifest_path = output_dir / "manifest.csv"
    old_paths = previous_manifest_paths(manifest_path)
    used_names: set[str] = set()
    manifest_rows: list[dict[str, str]] = []
    new_paths: set[Path] = set()

    for source in sources:
        relative = source.relative_to(source_dir)
        html_title, rows = parse_html(source)
        title = workbook_title(relative, html_title, section_headings)
        target_relative = output_relative_path(relative, title, used_names)
        target = output_dir / target_relative
        target.parent.mkdir(parents=True, exist_ok=True)
        make_workbook(relative, rows, title).save(target)
        new_paths.add(target_relative)
        manifest_rows.append({
            "source_html": relative.as_posix(),
            "xlsx_file": target_relative.as_posix(),
            "title": title,
        })
        print(f"{relative.as_posix()} -> {target_relative.as_posix()} ({len(rows)} элементов)")

    # Remove only files listed by the previous generated manifest; leave any
    # user-created workbooks in the output directory untouched.
    for old_relative in old_paths - new_paths:
        old_path = (output_dir / old_relative).resolve()
        if old_path.suffix.lower() == ".xlsx" and is_relative_to(old_path, output_dir):
            old_path.unlink(missing_ok=True)

    write_manifest(manifest_path, manifest_rows)
    return len(sources)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source-dir", type=Path, default=REPOSITORY_ROOT,
        help="Каталог с HTML-файлами (по умолчанию — корень репозитория).",
    )
    parser.add_argument(
        "--output-dir", type=Path, default=REPOSITORY_ROOT / "xlsx",
        help="Каталог для XLSX с сохранением структуры путей (по умолчанию — xlsx/).",
    )
    args = parser.parse_args(argv)

    try:
        count = convert(args.source_dir, args.output_dir)
    except (OSError, RuntimeError, ValueError) as error:
        print(f"Ошибка: {error}", file=sys.stderr)
        return 1
    print(f"Готово: создано файлов XLSX — {count}; каталог: {args.output_dir.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
